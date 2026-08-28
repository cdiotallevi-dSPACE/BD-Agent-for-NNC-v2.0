from __future__ import annotations
import json, shutil, tempfile
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from .core import normalize_company_domain

HEADERS={
    "company_id":"company_id", "company_name":"company_name",
    "other_names":"Other_names",
    "validated_source_assignees":"Validated Source assignees",
    "company_domain":"company_domain", "country":"country",
    "industry":"industry", "notes":"notes", "scan_flag":"Scan Y/N",
    "overall_rating":"Overall Product Match",
    "total_retrieved_sources":"Total Retrieved Sources",
    "total_retrieved_and_approved_sources":"Total Retrieved and approved Sources",
    "total_retrieved_pdfs":"total retrieved PDFs",
    "total_retrieved_and_approved_pdfs":"total retrieved and approved PDFs",
    "total_retrieved_patents":"total retrieved Patents",
    "total_retrieved_and_approved_patents":"total retrieved and approved patents",
}
COUNT_KEYS=tuple(key for key in HEADERS if key.startswith("total_"))
EXPECTED_COLUMNS={"Other_names":3,"Scan Y/N":8,"Overall Product Match":9,
    "Total Retrieved Sources":10,"Total Retrieved and approved Sources":11,
    "total retrieved PDFs":12,"total retrieved and approved PDFs":13,
    "total retrieved Patents":14,"total retrieved and approved patents":15,
    "Validated Source assignees":16}
def parse_other_names(value) -> list[str]:
    names=[]; seen=set()
    for part in str(value or "").split(","):
        name=part.strip(); key=name.casefold()
        if name and key not in seen: names.append(name); seen.add(key)
    return names
def parse_validated_source_assignees(value) -> list[str]:
    """Parse operator-approved legal identities; semicolon is preferred."""
    text=str(value or "").strip()
    if not text: return []
    parts=text.split(";") if ";" in text else text.split(",")
    names=[]; seen=set()
    for part in parts:
        name=part.strip(); key=name.casefold()
        if name and key not in seen: names.append(name); seen.add(key)
    return names
def normalize_previous_rating(value) -> int | float | None:
    if value is None or (isinstance(value,str) and not value.strip()): return None
    if isinstance(value,bool): raise ValueError("Overall Product Match must be blank or numeric")
    try: rating=float(value)
    except (TypeError,ValueError) as exc: raise ValueError("Overall Product Match must be blank, whitespace, or numeric") from exc
    if not 0<=rating<=100: raise ValueError("Overall Product Match must be between 0 and 100")
    return int(rating) if rating.is_integer() else rating
def selected_companies(path: Path, cycle_id: str) -> list[dict]:
    wb=load_workbook(path,read_only=True,data_only=False); ws=wb["Companies"]
    if ws.max_row is None:
        ws.calculate_dimension(force=True)
    headers={str(c.value).strip():c.column for c in ws[1] if c.value}
    missing=set(HEADERS.values())-headers.keys()
    if missing: raise ValueError(f"missing headers: {sorted(missing)}")
    misplaced={name:(headers.get(name),column) for name,column in EXPECTED_COLUMNS.items() if headers.get(name)!=column}
    if misplaced: raise ValueError(f"registry headers are in unexpected columns: {misplaced}")
    rows=[]
    for row in range(2,ws.max_row+1):
        flag=str(ws.cell(row,8).value or "").strip().upper(); raw_rating=ws.cell(row,9).value
        if flag=="Y":
            raw_domain=ws.cell(row,headers[HEADERS["company_domain"]]).value
            other_names_raw=ws.cell(row,headers[HEADERS["other_names"]]).value
            source_assignees_raw=ws.cell(row,headers[HEADERS["validated_source_assignees"]]).value
            rows.append({"cycle_id":cycle_id,"row_number":row,"company_id":ws.cell(row,headers[HEADERS["company_id"]]).value,"company_name":ws.cell(row,headers[HEADERS["company_name"]]).value,"other_names":parse_other_names(other_names_raw),"other_names_raw":other_names_raw,"validated_source_assignees":parse_validated_source_assignees(source_assignees_raw),"validated_source_assignees_raw":source_assignees_raw,"company_domain":normalize_company_domain(raw_domain),"company_domain_raw":raw_domain,"country":ws.cell(row,headers[HEADERS["country"]]).value,"industry":ws.cell(row,headers[HEADERS["industry"]]).value,"notes":ws.cell(row,headers[HEADERS["notes"]]).value,"scan_enabled":True,"previous_overall_product_match":normalize_previous_rating(raw_rating),"previous_overall_product_match_raw":raw_rating,"previous_cycle_source_counts_raw":{key:ws.cell(row,headers[HEADERS[key]]).value for key in COUNT_KEYS}})
    wb.close(); return rows
def update_rating(path: Path, company_id: str, company_name: str, previous_rating, rating: int, cycle_id: str, run_id: str, source_counts:dict|None=None, previous_source_counts:dict|None=None) -> dict:
    if not 0<=rating<=100: raise ValueError("rating must be 0..100")
    source_counts=source_counts or {key:0 for key in COUNT_KEYS}
    if set(source_counts)!=set(COUNT_KEYS): raise ValueError("source count fields are incomplete")
    if any(isinstance(value,bool) or not isinstance(value,int) or value<0 for value in source_counts.values()): raise ValueError("source counts must be non-negative integers")
    cycle_stamp=cycle_id.removeprefix("cycle_")
    backup=path.parent/"backups"/f"Companies_{cycle_stamp}.xlsx"; backup.parent.mkdir(exist_ok=True)
    if not backup.exists(): shutil.copy2(path,backup)
    wb=load_workbook(path); ws=wb["Companies"]; headers={str(c.value).strip():c.column for c in ws[1] if c.value}
    matches=[r for r in range(2,ws.max_row+1) if str(ws.cell(r,headers[HEADERS["company_id"]]).value)==str(company_id)]
    if len(matches)!=1 or str(ws.cell(matches[0],headers[HEADERS["company_name"]]).value)!=str(company_name): raise ValueError("company identity mismatch")
    rating_column=headers[HEADERS["overall_rating"]]
    current=ws.cell(matches[0],rating_column).value
    if current != previous_rating: raise ValueError(f"rating changed after cycle snapshot: expected {previous_rating!r}, found {current!r}")
    if previous_source_counts is not None:
        changed={key:(previous_source_counts.get(key),ws.cell(matches[0],headers[HEADERS[key]]).value) for key in COUNT_KEYS if ws.cell(matches[0],headers[HEADERS[key]]).value!=previous_source_counts.get(key)}
        if changed: raise ValueError(f"source counts changed after cycle snapshot: {changed}")
    ws.cell(matches[0],rating_column).value=rating
    for key,value in source_counts.items():
        cell=ws.cell(matches[0],headers[HEADERS[key]]); cell.value=value; cell.number_format='#,##0'
    with tempfile.NamedTemporaryFile(suffix=".xlsx",dir=path.parent,delete=False) as f: tmp=Path(f.name)
    try: wb.save(tmp); load_workbook(tmp,read_only=True).close(); tmp.replace(path)
    finally:
        if tmp.exists(): tmp.unlink()
    normalized_previous=normalize_previous_rating(previous_rating)
    return {"event":"company_cycle_results_written","cycle_id":cycle_id,"company_id":company_id,"company_name":company_name,"worksheet":"Companies","row":matches[0],"rating_update":{"header":"Overall Product Match","expected_column":"I","previous_value":previous_rating,"previous_numeric_value":normalized_previous,"new_value":rating,"score_change":None if normalized_previous is None else rating-normalized_previous},"source_count_updates":{key:{"header":HEADERS[key],"previous_value":(previous_source_counts or {}).get(key),"new_value":source_counts[key]} for key in COUNT_KEYS},"assessment_run_id":run_id,"timestamp":datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}

def queue_pending(path: Path, record: dict) -> None:
    path.parent.mkdir(parents=True,exist_ok=True)
    data=json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"updates":[]}
    for old in data["updates"]:
        if old["company_id"]==record["company_id"] and old["write_status"].startswith("pending"):
            old["write_status"]="superseded"
            old["superseded_by_assessment_run_id"]=record["assessment_run_id"]
    data["updates"].append(record)
    path.write_text(json.dumps(data,indent=2,sort_keys=True),encoding="utf-8")
